"""10_beautify_tables.py — 美化「阶段1/2 数据根」下的 .xlsx 表格（阶段1 收尾步）:
 通用: 表头(深蓝填充+白字加粗+居中+自动换行), 冻结首行, 按采样自动列宽
 小表(max_row<=1500): 斑马纹, 统一Arial字体, 数值列数字格式(整数0/小数0.0000)
 大表(拆分/公式表): 仅通用排版, 避免逐格样式拖慢
输入/输出: P.QUANT（原地美化 02_模型量化数据_GH1 根目录下的表）
"""
import os
import sys
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from pathlib import Path

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import gyo_paths as P          # 统一路径入口

D = Path(P.QUANT)
files = sorted([p for p in D.glob('*.xlsx') if not p.name.startswith('~$')])

HEADER_FILL = PatternFill('solid', fgColor='305496')
HEADER_FONT = Font(name='Arial', bold=True, color='FFFFFF', size=11)
ZEBRA = PatternFill('solid', fgColor='F2F6FB')
THIN = Side(style='thin', color='D9D9D9')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
DATA_FONT = Font(name='Arial', size=10)

for f in files:
    wb = load_workbook(f)
    for ws in wb.worksheets:
        nrows, ncols = ws.max_row, ws.max_column
        if nrows == 0 or ncols == 0:
            continue
        for c in range(1, ncols + 1):
            cell = ws.cell(row=1, column=c)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.border = BORDER
        ws.row_dimensions[1].height = 26
        ws.freeze_panes = 'A2'
        sample = min(nrows, 200)
        for c in range(1, ncols + 1):
            col = get_column_letter(c)
            maxlen = len(str(ws.cell(row=1, column=c).value or ''))
            for r in range(2, sample + 1):
                v = ws.cell(row=r, column=c).value
                if v is not None:
                    maxlen = max(maxlen, len(str(v)))
            ws.column_dimensions[col].width = min(max(maxlen + 2, 9), 44)
        if nrows <= 1500:
            for r in range(2, nrows + 1):
                for c in range(1, ncols + 1):
                    cell = ws.cell(row=r, column=c)
                    if r % 2 == 0:
                        cell.fill = ZEBRA
                    cell.font = DATA_FONT
                    v = cell.value
                    if isinstance(v, float):
                        cell.number_format = '0.0000'
                    elif isinstance(v, int):
                        cell.number_format = '0'
    wb.save(f)
    print('美化完成:', f.name, f'({len(wb.sheetnames)} sheets)')
print('全部完成。')
