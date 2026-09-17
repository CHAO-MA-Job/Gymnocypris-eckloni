"""0_2.数据表公式.py —— 【旧版·Excel 公式实现】运动学公式表（参考用）

⚠ 状态：**功能已被 `5_process_quanshen.py` 取代**（后者直接算数值，不再写 Excel 公式）。
   本脚本保留作**公式口径参考**；其列布局假设为旧版逐视频表
   （A=帧号, D=帧间dt, G–J=bbox, K/L=中心点, N=换算比例），
   与现行 `1_全身.xlsx`（A=帧号, B=时间戳, C=时间S, D=类别, …）**不一致**，
   故不做默认路径替换——必须显式指定输入，避免误用产生错误结果。

用法：set GYO_FORMULA_IN=<旧版 experimental data.xlsx>
      set GYO_FORMULA_OUT=<输出 处理后.xlsx>
      python 0_2.数据表公式.py
"""
import os
import sys
import openpyxl

# 文件路径（必须显式提供；旧版路径已失效，不设默认值）
file_path = os.environ.get('GYO_FORMULA_IN')
if not file_path:
    raise SystemExit(
        '【0_2.数据表公式.py】为旧版 Excel 公式实现，已被 5_process_quanshen.py 取代。\n'
        '若确需复现，请设 GYO_FORMULA_IN（旧版逐视频表）与 GYO_FORMULA_OUT（输出路径）。\n'
        '现行链路请直接运行 5_process_quanshen.py。')

# 打开Excel文件
wb = openpyxl.load_workbook(file_path)

def process_sheet(ws):
    """对单个sheet执行处理逻辑"""

    # === 操作1: 设置列公式 ===
    for row in range(3, ws.max_row + 1):
        ws[f"D{row}"] = f"=(A{row}-A{row-1})/30"
    for row in range(2, ws.max_row + 1):
        ws[f"K{row}"] = f"=(G{row}+I{row})/2"
        ws[f"L{row}"] = f"=(H{row}+J{row})/2"
        ws[f"M{row}"] = f'="(" & TEXT(K{row}, "0.00") & ", " & TEXT(L{row}, "0.00") & ")"'

    # === 操作2: 添加常量 ===
    ws["N1"] = "换算比例"
    ws["N2"] = "0.03125"

    # O列
    for row in range(3, ws.max_row + 1):
        ws[f"O{row}"] = f"=((K{row}-K{row-1})^2+(L{row}-L{row-1})^2)^(1/2)/D{row}"
    for row in range(3, ws.max_row + 1):
        ws[f"P{row}"] = f"=O{row}*$N$2"

    # Q列
    ws["Q2"] = 0
    for row in range(3, ws.max_row + 1):
        ws[f"Q{row}"] = f"=(((K{row}-K{row-1})^2+(L{row}-L{row-1})^2)^(1/2))*$N$2"

    # === 表头 ===
    ws["M1"] = "坐标"
    ws["O1"] = "v/像素点"
    ws["P1"] = "v/cm"
    ws["Q1"] = "S/单次位移"
    ws["R1"] = "S/总位移"
    ws["S1"] = "time/10s"
    ws["T1"] = "v/10s"
    ws["U1"] = "time/10秒"
    ws["V1"] = "v/10秒"

    # === R列：累计总位移 ===
    ws["R2"] = 0
    for row in range(3, ws.max_row + 1):
        ws[f"R{row}"] = f"=Q{row}+R{row-1}"

    # === S列：每300帧合并一次，换算时间 ===
    frame_step = 300
    fps = 30
    start_row = 2
    max_row = ws.max_row
    for block_start in range(start_row, max_row + 1, frame_step):
        block_end = min(block_start + frame_step - 1, max_row)
        target_row = block_start
        formula_row = block_end
        ws.merge_cells(start_row=target_row, end_row=block_end, start_column=19, end_column=19)  # S列
        ws[f"S{target_row}"] = f"=A{formula_row}/{fps}"

    # === T列：每300帧合并一次，计算P列平均速度 ===
    ws["P2"] = 0
    for block_start in range(start_row, max_row + 1, frame_step):
        block_end = min(block_start + frame_step - 1, max_row)
        target_row = block_start
        ws.merge_cells(start_row=target_row, end_row=block_end, start_column=20, end_column=20)  # T列
        ws[f"T{target_row}"] = f"=AVERAGE(P{block_start}:P{block_end})"

    # === U/V 列：公式引用 S/T 列（保持动态更新） ===
    uv_row = 2
    for block_start in range(start_row, max_row + 1, frame_step):
        ws[f"U{uv_row}"] = f"=S{block_start}"
        ws[f"V{uv_row}"] = f"=T{block_start}"
        uv_row += 1


    # === 固定 M 列宽度 ===
    ws.column_dimensions["M"].width = 20


# === 遍历所有sheet ===
for sheet in wb.sheetnames:
    ws = wb[sheet]
    print(f"正在处理工作表: {sheet}")
    process_sheet(ws)

# 保存处理后的文件
output_path = os.environ.get('GYO_FORMULA_OUT') or (file_path + '.处理后.xlsx')
wb.save(output_path)

print("处理完成，保存为：", output_path)
